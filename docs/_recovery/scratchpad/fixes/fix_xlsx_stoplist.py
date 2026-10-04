import openpyxl
p = "/Users/dmitry/Projects/teta_new/docs/02_competitor_analysis/feature_matrix.xlsx"
wb = openpyxl.load_workbook(p)
ws = wb["Матрица функций"]
hdr_row = None
for r in range(1, 6):
    vals = [str(c.value or "") for c in ws[r]]
    if any("Функция" == v.strip() for v in vals):
        hdr_row = r; hdr = vals; break
col = {name.strip(): i + 1 for i, name in enumerate(hdr)}
fcol, dcol = col["Функция"], [v for k, v in col.items() if k.startswith("Решение")][0]
scol = [v for k, v in col.items() if k.startswith("Раздел")][0]
n = 0
for r in range(hdr_row + 1, ws.max_row + 1):
    v = str(ws.cell(r, fcol).value or "")
    if "Стоп-лист" in v:
        print("row", r, "before:", ws.cell(r, dcol).value, "|", ws.cell(r, scol).value)
        ws.cell(r, dcol).value = "MVP"
        ws.cell(r, scol).value = "ADM-08 (MVP), HR-02 (R2)"
        n += 1
wb.save(p); print("updated rows:", n)
